"""Admin Excel export of graduates, tickets and payments (FR-19, D17).

One sheet, right-to-left, Arabic headers, one row per graduate of the selected
event(s). A graduate without a ticket still appears, with empty ticket and
payment cells; an event without data yields a header-only file (UC-9 alt).

Safety (D17): every user-controlled text cell is written as an explicit string
(``data_type="s"``) so a value such as ``=HYPERLINK(...)`` can never be treated
as a formula by a spreadsheet program. The numeric columns (companions count,
price and amount) are written as real numbers so a spreadsheet can sum them;
they come from integer/decimal columns, never from free text.
"""

from datetime import UTC, datetime
from io import BytesIO

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import and_, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core import messages
from app.models.college import CollegeModel
from app.models.event import EventModel
from app.models.graduate import GraduateModel
from app.models.payment import PaymentModel
from app.models.ticket import TicketModel
from app.models.major import MajorModel
from app.services.audit_log import ActorType, record_audit

CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_HEADERS = (
    "اسم الحفل",
    "الكلية",
    "الاسم الكامل",
    "الجنس",
    "الرقم الجامعي",
    "اسم الأب",
    "اسم الأم",
    "البريد الإلكتروني",
    "الهاتف",
    "البريد موثّق",
    "نوع التذكرة",
    "عدد المرافقين",
    "السعر",
    "حالة التذكرة",
    "سبب الرفض",
    "رقم العملية",
    "المبلغ",
    "اسم المرسل",
    "تاريخ الإرسال",
    "تاريخ المراجعة",
    "التخصص",
    "المعدل",
    "الترتيب",
)


def _enum_value(value) -> str:
    """Return the raw value of an enum member, or ``""`` for None."""
    if value is None:
        return ""
    return str(getattr(value, "value", value))


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def _string_cell(ws, row: int, col: int, value) -> None:
    """Write a value as an explicit string so it can never be a formula (D17)."""
    cell = ws.cell(row=row, column=col)
    cell.value = _as_text(value)
    cell.data_type = "s"


def _number_cell(ws, row: int, col: int, value) -> None:
    """Write a numeric value as a real number so it can be summed.

    ``None`` (for example a graduate without a ticket) leaves the cell empty.
    Booleans are rejected to avoid ``True`` silently becoming ``1``.
    """
    if value is None or value == "" or isinstance(value, bool):
        return
    cell = ws.cell(row=row, column=col)
    cell.value = int(value) if isinstance(value, int) else float(value)
    cell.data_type = "n"


def _format_rank(rank) -> str:
    if rank == 1:
        return "الأول"
    if rank == 2:
        return "الثاني"
    if rank == 3:
        return "الثالث"
    return ""


async def export_service(db: AsyncSession, event_id: int | None, admin: dict):
    """Build the .xlsx export and return ``(content_bytes, filename)``."""
    event_name = None
    if event_id is not None:
        event = await db.get(EventModel, event_id)
        if event is None:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=messages.EVENT_NOT_FOUND,
            )
        event_name = event.name

    latest_payment = (
        select(PaymentModel.ticket_id, func.max(PaymentModel.id).label("payment_id"))
        .group_by(PaymentModel.ticket_id)
        .subquery()
    )

    stmt = (
        select(
            EventModel.name,
            CollegeModel.name,
            GraduateModel,
            TicketModel,
            PaymentModel,
            MajorModel.name.label("major_name"),
        )
        .join(CollegeModel, CollegeModel.id == GraduateModel.college_id)
        .join(EventModel, EventModel.id == CollegeModel.event_id)
        .outerjoin(MajorModel, MajorModel.id == GraduateModel.major_id)
        .outerjoin(
            TicketModel,
            and_(
                TicketModel.graduate_id == GraduateModel.id,
                TicketModel.event_id == EventModel.id,
            ),
        )
        .outerjoin(latest_payment, latest_payment.c.ticket_id == TicketModel.id)
        .outerjoin(PaymentModel, PaymentModel.id == latest_payment.c.payment_id)
        .order_by(
            EventModel.id.asc(), CollegeModel.id.asc(), GraduateModel.id.asc()
        )
    )
    if event_id is not None:
        stmt = stmt.where(EventModel.id == event_id)

    rows = (await db.execute(stmt)).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "المتخرجون"
    ws.sheet_view.rightToLeft = True

    for col_idx, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.data_type = "s"
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="right", vertical="center")

    for row_idx, (
        ev_name,
        college_name,
        graduate,
        ticket,
        payment,
        major_name,
    ) in enumerate(rows, start=2):
        _string_cell(ws, row_idx, 1, ev_name)
        _string_cell(ws, row_idx, 2, college_name)
        _string_cell(ws, row_idx, 3, graduate.full_name)
        _string_cell(ws, row_idx, 4, graduate.gender)
        _string_cell(ws, row_idx, 5, graduate.university_id)
        _string_cell(ws, row_idx, 6, graduate.father_name)
        _string_cell(ws, row_idx, 7, graduate.mother_name)
        _string_cell(ws, row_idx, 8, graduate.email)
        _string_cell(ws, row_idx, 9, graduate.phone)
        _string_cell(ws, row_idx, 10, "نعم" if graduate.email_verified else "لا")
        _string_cell(ws, row_idx, 11, _enum_value(ticket.ticket_type) if ticket else "")
        _number_cell(ws, row_idx, 12, ticket.companions_count if ticket else None)
        _number_cell(ws, row_idx, 13, ticket.price if ticket else None)
        _string_cell(ws, row_idx, 14, _enum_value(ticket.status) if ticket else "")
        _string_cell(ws, row_idx, 15, ticket.rejection_reason if ticket else "")
        _string_cell(ws, row_idx, 16, payment.transaction_number if payment else "")
        _number_cell(ws, row_idx, 17, payment.amount if payment else None)
        _string_cell(ws, row_idx, 18, payment.sender_name if payment else "")
        _string_cell(ws, row_idx, 19, payment.submitted_at if payment else "")
        _string_cell(ws, row_idx, 20, payment.reviewed_at if payment else "")
        _string_cell(ws, row_idx, 21, major_name)
        _string_cell(ws, row_idx, 22, str(graduate.gpa) if graduate.gpa is not None else "")
        _string_cell(ws, row_idx, 23, _format_rank(graduate.top_rank))

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    content = buffer.getvalue()

    scope = f"event_id={event_id}" if event_id is not None else "all events"
    target = f"Event ID {event_id}" if event_id is not None else "All Events"
    await record_audit(
        db,
        actor_type=ActorType.ADMIN,
        actor_id=admin["admin_id"],
        action="DATA_EXPORTED",
        target=target,
        details=f"scope={scope}; rows={len(rows)}",
        event_id=event_id,
    )

    day = datetime.now(UTC).strftime("%Y%m%d")
    scope_slug = str(event_id) if event_id is not None else "all"
    filename = f"graduation_export_{scope_slug}_{day}.xlsx"
    return content, filename
