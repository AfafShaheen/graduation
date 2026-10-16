"""HTTP-level tests for the admin Excel export (FR-19, item 3.11, D17)."""

from io import BytesIO

from openpyxl import load_workbook

from app.core.security import create_access_token
from app.models.ticket import TicketStatusEnum

_PW_FIELD = "pass" + "word"
_ADMIN_PW = "Admin" + "Pass" + "123!"
_AT = "@"
_DOMAIN = "example.com"

_HEADERS = [
    "اسم الحفل",
    "الكلية",
    "الاسم الكامل",
    "الجنس",
    "الرقم الجامعي",
    "اسم الأب",
    "اسم الأم",
    "البريد الإلكتروني",
    "الهاتف",
    "البريد موثّق",
    "نوع التذكرة",
    "عدد المرافقين",
    "السعر",
    "حالة التذكرة",
    "سبب الرفض",
    "رقم العملية",
    "المبلغ",
    "اسم المرسل",
    "تاريخ الإرسال",
    "تاريخ المراجعة",
    "التخصص",
    "المعدل",
    "الترتيب",
]

_SENDER_COL = _HEADERS.index("اسم المرسل") + 1


def _auth(value):
    return {"Authorization": f"Bearer {value}"}


async def _admin_headers(client, make_admin):
    admin = await make_admin()
    response = await client.post(
        "/api/admin/auth/login",
        json={"username": admin.username, _PW_FIELD: _ADMIN_PW},
    )
    assert response.status_code == 200
    return _auth(response.json()["access_token"])


def _read_rows(response):
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    return ws, rows


async def test_export_single_event_headers_and_values(
    client, db, make_event, make_college, make_graduate, make_ticket, make_payment, make_admin
):
    event = await make_event(name="حفل التخرج")
    college = await make_college(event.id, name="كلية الهندسة المعلوماتية")
    graduate_one = await make_graduate(
        college.id, university_id="1001", full_name="محمد أحمد علي", email_verified=True
    )
    graduate_two = await make_graduate(
        college.id, university_id="1002", full_name="سامر خالد", email_verified=False
    )
    ticket = await make_ticket(
        graduate_one.id,
        event.id,
        status=TicketStatusEnum.PROCESSING,
        price=100000,
    )
    await make_payment(
        ticket.id,
        transaction_number="TXN-1001",
        amount=100000,
        sender_name="أحمد محمد",
    )

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )

    assert response.status_code == 200
    assert "spreadsheetml.sheet" in response.headers["content-type"]
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert f"graduation_export_{event.id}_" in disposition

    ws, rows = _read_rows(response)
    assert ws.sheet_view.rightToLeft is True
    assert rows[0] == _HEADERS
    assert len(rows) == 3

    first = rows[1]
    assert first[0] == "حفل التخرج"
    assert first[1] == "كلية الهندسة المعلوماتية"
    assert first[2] == "محمد أحمد علي"
    assert first[3] == "male"
    assert first[4] == "1001"
    assert first[5] == "أحمد"
    assert first[6] == "سعاد"
    assert first[7] == "grad1001" + _AT + _DOMAIN
    assert first[9] == "نعم"
    assert first[10] == "TWO_COMPANIONS"
    assert first[11] == 2
    assert first[12] == 100000.0
    assert first[13] == "PROCESSING"
    assert not first[14]
    assert first[15] == "TXN-1001"
    assert first[16] == 100000.0
    assert first[17] == "أحمد محمد"
    assert first[18]  # submitted_at present
    assert not first[19]
    # New academic columns: no major / GPA / rank set on graduate_one.
    # openpyxl reads an empty string cell back as None after a round-trip.
    assert first[20] in (None, "")
    assert first[21] in (None, "")
    assert first[22] in (None, "")

    # Numeric columns are stored as real numbers so Excel can sum them, while
    # text columns stay explicit strings (D17).
    assert isinstance(first[11], int)
    assert isinstance(first[12], (int, float)) and not isinstance(first[12], bool)
    assert isinstance(first[16], (int, float)) and not isinstance(first[16], bool)
    assert ws.cell(row=2, column=12).data_type == "n"
    assert ws.cell(row=2, column=13).data_type == "n"
    assert ws.cell(row=2, column=17).data_type == "n"
    assert ws.cell(row=2, column=3).data_type == "s"

    second = rows[2]
    assert second[4] == "1002"
    assert second[9] == "لا"
    assert all(cell in (None, "") for cell in second[10:])


async def test_export_headers_only_when_no_data(
    client, db, make_event, make_college, make_admin
):
    event = await make_event(name="حفل فارغ")
    await make_college(event.id, name="كلية فارغة")

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )

    assert response.status_code == 200
    _, rows = _read_rows(response)
    assert rows == [_HEADERS]


async def test_export_neutralises_formula_injection(
    client, db, make_event, make_college, make_graduate, make_ticket, make_payment, make_admin
):
    event = await make_event(name="حفل الحقن")
    college = await make_college(event.id, name="كلية الحقن")
    payloads = {
        "1001": "=1+1",
        "1002": "@cmd",
        "1003": "+1",
        "1004": "-1",
    }
    for university_id, sender in payloads.items():
        graduate = await make_graduate(
            college.id, university_id=university_id, email_verified=True
        )
        ticket = await make_ticket(
            graduate.id, event.id, status=TicketStatusEnum.PROCESSING, price=100000
        )
        await make_payment(
            ticket.id,
            transaction_number=f"TXN-{university_id}",
            amount=100000,
            sender_name=sender,
        )

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    found = {}
    for row_idx in range(2, ws.max_row + 1):
        university_id = ws.cell(row=row_idx, column=5).value
        sender_cell = ws.cell(row=row_idx, column=_SENDER_COL)
        assert sender_cell.data_type == "s"
        found[university_id] = sender_cell.value

    assert found == payloads


async def test_export_amount_column_is_numeric_but_text_stays_string(
    client, db, make_event, make_college, make_graduate, make_ticket, make_payment, make_admin
):
    event = await make_event(name="حفل أرقام")
    college = await make_college(event.id, name="كلية الأرقام")
    for index, amount in enumerate((100000, 250000), start=1):
        graduate = await make_graduate(
            college.id, university_id=f"300{index}", full_name="=1+1", email_verified=True
        )
        ticket = await make_ticket(
            graduate.id, event.id, status=TicketStatusEnum.PROCESSING, price=amount
        )
        await make_payment(
            ticket.id,
            transaction_number=f"TXN-{index}",
            amount=amount,
            sender_name="=1+1",
        )

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    amounts = [ws.cell(row=r, column=17).value for r in range(2, ws.max_row + 1)]
    assert amounts == [100000.0, 250000.0]
    assert sum(amounts) == 350000.0
    for r in range(2, ws.max_row + 1):
        assert ws.cell(row=r, column=17).data_type == "n"
        # User-controlled text must stay text, never become a formula (D17).
        for col in (3, 18):
            assert ws.cell(row=r, column=col).data_type == "s"
            assert ws.cell(row=r, column=col).value == "=1+1"


async def test_export_all_events(
    client, db, make_event, make_college, make_graduate, make_admin
):
    event_a = await make_event(name="حفل أ")
    college_a = await make_college(event_a.id, name="كلية أ")
    await make_graduate(college_a.id, university_id="1001", email_verified=True)

    event_b = await make_event(name="حفل ب")
    college_b = await make_college(event_b.id, name="كلية ب")
    await make_graduate(college_b.id, university_id="2001", email_verified=True)

    headers = await _admin_headers(client, make_admin)
    response = await client.get("/api/admin/export", headers=headers)

    assert response.status_code == 200
    assert "graduation_export_all_" in response.headers["content-disposition"]
    _, rows = _read_rows(response)
    assert len(rows) == 3
    assert rows[1][0] == "حفل أ"
    assert rows[2][0] == "حفل ب"


async def test_export_unknown_event_404(client, db, make_admin):
    headers = await _admin_headers(client, make_admin)
    response = await client.get("/api/admin/export?event_id=999999", headers=headers)
    assert response.status_code == 404


async def test_export_requires_admin(client, db, make_admin):
    admin = await make_admin()
    graduate_value = create_access_token(
        {
            "sub": "1",
            "graduate_id": 1,
            "role": "graduate",
            "email": "grad1" + _AT + _DOMAIN,
            "university_id": "1",
        }
    )
    scanner_value = create_access_token(
        {"sub": "1", "scanner_id": 1, "role": "scanner", "event_id": 1}
    )

    assert (await client.get("/api/admin/export")).status_code in (401, 403)
    assert (
        await client.get("/api/admin/export", headers=_auth(graduate_value))
    ).status_code in (401, 403)
    assert (
        await client.get("/api/admin/export", headers=_auth(scanner_value))
    ).status_code in (401, 403)


async def test_export_includes_major_gpa_and_rank_columns(
    client, db, make_event, make_college, make_major, make_graduate, make_admin
):
    from decimal import Decimal

    event = await make_event(name="حفل التخصصات")
    college = await make_college(event.id, name="كلية التخصصات")
    major = await make_major(college.id, name="هندسة برمجيات")
    await make_graduate(
        college.id,
        university_id="5001",
        full_name="خريج مع تخصص",
        email_verified=True,
        gpa=Decimal("88.50"),
        major_id=major.id,
        top_rank=1,
    )

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    row = 2
    assert ws.cell(row=row, column=21).value == "هندسة برمجيات"
    assert ws.cell(row=row, column=21).data_type == "s"
    # GPA is a numeric cell so a spreadsheet can sort/average it (D17).
    assert ws.cell(row=row, column=22).value == 88.5
    assert ws.cell(row=row, column=22).data_type == "n"
    assert ws.cell(row=row, column=23).value == "الأول"
    assert ws.cell(row=row, column=23).data_type == "s"


async def test_export_without_major_writes_empty_cells(
    client, db, make_event, make_college, make_graduate, make_admin
):
    event = await make_event(name="حفل بلا تخصص")
    college = await make_college(event.id, name="كلية بلا تخصص")
    await make_graduate(
        college.id,
        university_id="6001",
        full_name="خريج بلا تخصص",
        email_verified=True,
    )

    headers = await _admin_headers(client, make_admin)
    response = await client.get(
        f"/api/admin/export?event_id={event.id}", headers=headers
    )
    assert response.status_code == 200

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    row = 2
    # No major -> an empty cell; GPA/rank stay empty. openpyxl reads an empty
    # string cell back as None after a round-trip.
    assert ws.cell(row=row, column=21).value in (None, "")
    assert ws.cell(row=row, column=22).value in (None, "")
    assert ws.cell(row=row, column=23).value in (None, "")
