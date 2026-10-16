"""Tests for admin filters (top-3 rank, volunteers, event_id, major_id, etc.)."""

import pytest
from decimal import Decimal
from openpyxl import load_workbook
from io import BytesIO

_ADMIN_PW = "Admin" + "Pass" + "123!"
_GRAD_PW = "Grad" + "Pass" + "1234"


async def _admin_headers(client, make_admin, username="admin_filters"):
    admin = await make_admin(username=username)
    response = await client.post(
        "/api/admin/auth/login",
        json={"username": admin.username, "password": _ADMIN_PW},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def _grad_headers(client, graduate):
    response = await client.post(
        "/api/auth/graduate/login",
        json={"email": graduate.email, "password": _GRAD_PW},
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def test_volunteers_top_rank_filters_and_stats_independence(client, db, make_admin, make_event, make_college, make_graduate):
    event = await make_event()
    college = await make_college(event.id)

    g1 = await make_graduate(college.id, university_id="1", top_rank=1, email_verified=True)
    g2 = await make_graduate(college.id, university_id="2", top_rank=2, email_verified=True)
    g3 = await make_graduate(college.id, university_id="3", top_rank=None, email_verified=True)

    from app.models.volunteer_answer import VolunteerAnswerModel
    db.add(VolunteerAnswerModel(graduate_id=g1.id, wants_responsibility=True, has_experience=False))
    db.add(VolunteerAnswerModel(graduate_id=g2.id, wants_responsibility=False, has_experience=True))
    db.add(VolunteerAnswerModel(graduate_id=g3.id, wants_responsibility=True, has_experience=True))
    await db.commit()

    headers = await _admin_headers(client, make_admin, username="admin_vol1")

    # 1. Test top_rank filter on /api/admin/volunteers
    r1 = await client.get("/api/admin/volunteers?top_rank=1", headers=headers)
    assert r1.status_code == 200
    assert r1.json()["total"] == 1
    assert r1.json()["items"][0]["graduate_id"] == g1.id

    r_any = await client.get("/api/admin/volunteers?top_rank=any", headers=headers)
    assert r_any.json()["total"] == 2

    r_none = await client.get("/api/admin/volunteers?top_rank=none", headers=headers)
    assert r_none.json()["total"] == 1
    assert r_none.json()["items"][0]["graduate_id"] == g3.id

    r_invalid = await client.get("/api/admin/volunteers?top_rank=4", headers=headers)
    assert r_invalid.status_code == 422

    # 2. Stats stay independent of top_rank filter
    r_stats = await client.get("/api/admin/volunteers?top_rank=1", headers=headers)
    assert r_stats.json()["stats"]["all"] == 3


async def test_admin_graduates_filters_and_cache(client, db, make_admin, make_event, make_college, make_major, make_graduate):
    event = await make_event()
    college = await make_college(event.id)
    major = await make_major(college.id, name="حاسبات")

    g1 = await make_graduate(college.id, university_id="101", top_rank=1, gpa=Decimal("95.00"), major_id=major.id, email_verified=True)
    g2 = await make_graduate(college.id, university_id="102", top_rank=None, gpa=Decimal("75.00"), major_id=major.id, email_verified=True)

    from app.models.volunteer_answer import VolunteerAnswerModel
    db.add(VolunteerAnswerModel(graduate_id=g1.id, wants_responsibility=True, has_experience=False))
    await db.commit()

    headers = await _admin_headers(client, make_admin, username="admin_grads1")

    # top_rank=1
    resp = await client.get("/api/admin/graduates?top_rank=1", headers=headers)
    assert resp.status_code == 200
    items = resp.json()
    assert len(items) == 1
    assert items[0]["id"] == g1.id
    assert items[0]["major_name"] == "حاسبات"
    assert items[0]["gpa"] == "95.00"
    assert items[0]["top_rank"] == 1
    assert items[0]["wants_responsibility"] == "yes"
    assert items[0]["volunteer_answered"] is True

    # Cache test: call without filters, then with top_rank=1
    all_resp = await client.get("/api/admin/graduates", headers=headers)
    assert len(all_resp.json()) == 2
    top1_resp = await client.get("/api/admin/graduates?top_rank=1", headers=headers)
    assert len(top1_resp.json()) == 1

    # volunteer_answered=false
    none_resp = await client.get("/api/admin/graduates?volunteer_answered=false", headers=headers)
    assert len(none_resp.json()) == 1
    assert none_resp.json()[0]["id"] == g2.id

    # wants_responsibility=no should not return un-answered g2
    no_resp = await client.get("/api/admin/graduates?wants_responsibility=no", headers=headers)
    assert len(no_resp.json()) == 0

    # Unknown event_id -> 404
    bad_event = await client.get("/api/admin/graduates?event_id=999999", headers=headers)
    assert bad_event.status_code == 404

    # No token -> 401/403
    unauth = await client.get("/api/admin/graduates")
    assert unauth.status_code in (401, 403)


async def test_graduate_volunteer_put_invalidates_admin_cache(client, db, make_admin, make_event, make_college, make_graduate):
    admin = await make_admin(username="admin_cache")
    event = await make_event()
    college = await make_college(event.id)
    g = await make_graduate(college.id, university_id="200", email_verified=True)

    headers = await _admin_headers(client, make_admin, username="admin_cache2")
    # Initially 0 answered
    resp1 = await client.get("/api/admin/graduates?volunteer_answered=true", headers=headers)
    assert len(resp1.json()) == 0

    # Graduate answers via PUT
    grad_headers = await _grad_headers(client, g)
    await client.put("/api/graduates/me/volunteer-answers", json={"wants_responsibility": "yes", "has_experience": "yes"}, headers=grad_headers)

    # Admin list should reflect immediately due to cache invalidation
    resp2 = await client.get("/api/admin/graduates?volunteer_answered=true", headers=headers)
    assert len(resp2.json()) == 1
    assert resp2.json()[0]["id"] == g.id


async def test_volunteers_export_with_academic_columns(client, db, make_admin, make_event, make_college, make_major, make_graduate):
    admin = await make_admin(username="admin_export")
    event = await make_event()
    college = await make_college(event.id)
    major = await make_major(college.id, name="ذكاء اصطناعي")
    g = await make_graduate(college.id, university_id="300", top_rank=2, gpa=Decimal("92.50"), major_id=major.id, email_verified=True)

    from app.models.volunteer_answer import VolunteerAnswerModel
    db.add(VolunteerAnswerModel(graduate_id=g.id, wants_responsibility=True, has_experience=True))
    await db.commit()

    headers = await _admin_headers(client, make_admin, username="admin_export2")
    export_resp = await client.get("/api/admin/volunteers/export?top_rank=any", headers=headers)
    assert export_resp.status_code == 200

    wb = load_workbook(BytesIO(export_resp.content))
    ws = wb.active
    assert ws.cell(row=1, column=8).value == "التخصص"
    assert ws.cell(row=1, column=9).value == "المعدل"
    assert ws.cell(row=1, column=10).value == "الترتيب"
    assert ws.cell(row=2, column=8).value == "ذكاء اصطناعي"
    assert ws.cell(row=2, column=9).value == "92.50"
    assert ws.cell(row=2, column=10).value == "الثاني"
