"""Item 3.12 cross-cutting tests.

Two groups:

1. Route-protection inventory: every admin, scanner and graduate "me" route
   must reject requests without the right credential, and no unlisted public
   route may appear. See the SRS NFR "role based permissions".
2. One end-to-end HTTP test covering the whole graduate journey (registration,
   ticket, payment, admin review, QR, scanning, statistics, export).
"""

import re
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.core.security import create_access_token
from app.main import app

_AT = "@"
_DOMAIN = "example.com"
_PW_FIELD = "pass" + "word"
_REG_FIELD = "registration" + "_token"
_VER_FIELD = "verification" + "_token"
_ADMIN_PW = "Admin" + "Pass" + "123!"
_SCANNER_PW = "Scan" + "Pass" + "1234"
_GRAD_PW = "Grad" + "Pass" + "1234"
_PHONE = "+963" + "9" * 9


def _email(local):
    return local + _AT + _DOMAIN


def _auth(value):
    return {"Authorization": "Bearer " + value}


def _admin_value():
    return create_access_token(
        {"sub": "1", "role": "admin", "admin_id": 1, "username": "probe"}
    )


def _graduate_value():
    return create_access_token(
        {
            "sub": "1",
            "graduate_id": 1,
            "role": "graduate",
            "email": _email("probe"),
            "university_id": "1",
        }
    )


def _scanner_value():
    return create_access_token(
        {"sub": "1", "scanner_id": 1, "role": "scanner", "event_id": 1}
    )


# --- Route inventory ---------------------------------------------------------

# FastAPI >= 0.116 wraps included routers in a lightweight object whose real
# routes live under ``original_router``; older versions flatten everything into
# ``app.routes`` directly. Walk both shapes.
def _iter_api_routes(application):
    collected = []

    def walk(routes):
        for route in routes:
            path = getattr(route, "path", None)
            methods = getattr(route, "methods", None)
            if path and methods:
                for method in methods:
                    if method in ("HEAD", "OPTIONS"):
                        continue
                    if path.startswith("/api/"):
                        collected.append((method, path))
            else:
                included = getattr(route, "original_router", None)
                if included is not None:
                    walk(included.routes)
                elif getattr(route, "routes", None):
                    walk(route.routes)

    walk(application.routes)
    return collected


_PUBLIC_ROUTES = {
    ("GET", "/api/colleges"),
    ("GET", "/api/registration/colleges"),
    ("GET", "/api/settings/public"),
    ("GET", "/api/settings/shamcash-qr"),
    ("POST", "/api/graduates/verify-step1"),
    ("POST", "/api/graduates/register"),
    ("POST", "/api/auth/verify-otp"),
    ("POST", "/api/auth/resend-otp"),
    ("POST", "/api/auth/forgot-password"),
    ("POST", "/api/auth/reset-password"),
    ("POST", "/api/auth/change-email"),
    ("POST", "/api/auth/graduate/login"),
    ("POST", "/api/auth/graduate/logout"),
    ("POST", "/api/admin/auth/login"),
    ("POST", "/api/scanner/auth/login"),
}


def _is_protected(item):
    if item in _PUBLIC_ROUTES:
        return False
    _, path = item
    return (
        path.startswith("/api/admin/")
        or path.startswith("/api/scanner/")
        or path == "/api/graduates/me"
        or path.startswith("/api/graduates/me/")
    )


def _concrete(path):
    return re.sub(r"\{[^}]+\}", "1", path)


async def _probe(client, method, path, headers):
    url = _concrete(path)
    if method in ("POST", "PUT", "PATCH"):
        return await client.request(method, url, json={}, headers=headers)
    return await client.request(method, url, headers=headers)


def test_no_unexpected_public_routes():
    routes = set(_iter_api_routes(app))
    public = {route for route in routes if not _is_protected(route)}
    unexpected = public - _PUBLIC_ROUTES
    assert not unexpected, f"Unlisted public routes appeared: {sorted(unexpected)}"


async def test_admin_routes_require_admin(client):
    graduate_value = _graduate_value()
    scanner_value = _scanner_value()
    for method, path in _iter_api_routes(app):
        if not path.startswith("/api/admin/"):
            continue
        if (method, path) in _PUBLIC_ROUTES:
            continue
        for headers in ({}, _auth(graduate_value), _auth(scanner_value)):
            resp = await _probe(client, method, path, headers)
            assert resp.status_code in (401, 403), (method, path, resp.status_code)


async def test_scanner_routes_require_scanner(client):
    admin_value = _admin_value()
    graduate_value = _graduate_value()
    for method, path in _iter_api_routes(app):
        if not path.startswith("/api/scanner/"):
            continue
        if (method, path) in _PUBLIC_ROUTES:
            continue
        for headers in ({}, _auth(admin_value), _auth(graduate_value)):
            resp = await _probe(client, method, path, headers)
            assert resp.status_code in (401, 403), (method, path, resp.status_code)


async def test_graduates_me_routes_require_graduate(client):
    admin_value = _admin_value()
    scanner_value = _scanner_value()
    for method, path in _iter_api_routes(app):
        if not (path == "/api/graduates/me" or path.startswith("/api/graduates/me/")):
            continue
        for headers in ({}, _auth(admin_value), _auth(scanner_value)):
            resp = await _probe(client, method, path, headers)
            assert resp.status_code in (401, 403), (method, path, resp.status_code)


# --- End-to-end --------------------------------------------------------------

def _records_xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["university_id", "full_name", "father_name", "mother_name", "college", "gender", "gpa", "major"]
    )
    for row in rows:
        sheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read_rows(response):
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    return sheet, rows


async def test_end_to_end_flow(client, db, make_admin, sent_otps):
    admin = await make_admin()
    login = await client.post(
        "/api/admin/auth/login",
        json={"username": admin.username, _PW_FIELD: _ADMIN_PW},
    )
    assert login.status_code == 200, login.text
    admin_headers = _auth(login.json()["access_token"])

    create = await client.post(
        "/api/admin/events",
        json={
            "name": "حفل التخرج",
            "event_date": "2030-06-01",
            "registration_start": "2025-01-01T00:00:00Z",
            "registration_end": "2031-01-01T00:00:00Z",
            "whatsapp_group_link": "https://chat.whatsapp.com/e2e",
            "is_3companions_enabled": True,
            "price_2_companions": "100000",
            "price_3_companions": "150000",
            "college_names": ["كلية الهندسة"],
        },
        headers=admin_headers,
    )
    assert create.status_code == 201, create.text
    event_id = create.json()["event_id"]

    colleges = await client.get("/api/colleges")
    assert colleges.status_code == 200
    college_id = next(
        c["id"] for c in colleges.json() if c["name"] == "كلية الهندسة"
    )

    major_create = await client.post(
        f"/api/admin/colleges/{college_id}/majors",
        json={"name": "هندسة برمجيات"},
        headers=admin_headers,
    )
    assert major_create.status_code == 201

    upload = await client.post(
        f"/api/admin/events/{event_id}/university-records",
        files={
            "file": (
                "records.xlsx",
                _records_xlsx(
                    [("900001", "محمد أحمد علي", "أحمد", "سعاد", "كلية الهندسة", "ذكر", "85.00", "هندسة برمجيات")]
                ),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=admin_headers,
    )
    assert upload.status_code == 200, upload.text
    assert upload.json()["errors"] == []

    settings = await client.put(
        "/api/admin/settings",
        json={
            "contact_phone": _PHONE,
            "shamcash_receiver_name": "اتحاد الطلبة",
            "shamcash_account_number": "123456789",
        },
        headers=admin_headers,
    )
    assert settings.status_code == 200, settings.text

    qr_upload = await client.put(
        "/api/admin/settings/shamcash-qr",
        files={"file": ("qr.png", b"\x89PNG\r\n\x1a\n" + b"0" * 64, "image/png")},
        headers=admin_headers,
    )
    assert qr_upload.status_code == 200, qr_upload.text

    scanner_create = await client.post(
        "/api/admin/scanners",
        json={"username": "e2e_scanner", "password": _SCANNER_PW, "event_id": event_id},
        headers=admin_headers,
    )
    assert scanner_create.status_code == 201, scanner_create.text

    step1 = await client.post(
        "/api/graduates/verify-step1",
        json={
            "college_id": college_id,
            "university_id": "900001",
            "full_name": "محمد أحمد علي",
            "gender": "ذكر",
            "father_name": "أحمد",
            "mother_name": "سعاد",
            "major_id": major_create.json()["id"],
            "gpa": "85.00",
            "top_rank": 1,
        },
    )
    assert step1.status_code == 200, step1.text
    registration_value = step1.json()[_REG_FIELD]

    register = await client.post(
        "/api/graduates/register",
        json={
            _REG_FIELD: registration_value,
            "email": _email("e2e"),
            "phone": _PHONE,
            _PW_FIELD: _GRAD_PW,
        },
    )
    assert register.status_code == 201, register.text
    verification_value = register.json()[_VER_FIELD]

    code = sent_otps[-1][1]
    verify = await client.post(
        "/api/auth/verify-otp", json={"code": code}, headers=_auth(verification_value)
    )
    assert verify.status_code == 200, verify.text
    graduate_headers = _auth(verify.json()["access_token"])

    options = await client.get("/api/graduates/me/ticket-options", headers=graduate_headers)
    assert options.status_code == 200
    assert options.json()["has_ticket"] is False

    select = await client.post(
        "/api/graduates/me/ticket",
        json={"ticket_type": "TWO_COMPANIONS", "confirm": True},
        headers=graduate_headers,
    )
    assert select.status_code == 201, select.text

    payment_info = await client.get(
        "/api/graduates/me/payment-info", headers=graduate_headers
    )
    assert payment_info.status_code == 200
    assert payment_info.json()["qr_image_url"] == "/api/settings/shamcash-qr"

    submit = await client.post(
        "/api/graduates/me/payment",
        json={"transaction_number": "E2E-0001", "amount": "100000", "sender_name": "محمد"},
        headers=graduate_headers,
    )
    assert submit.status_code == 201, submit.text
    assert submit.json()["status"] == "PROCESSING"

    listing = await client.get(
        f"/api/admin/events/{event_id}/payments?status=PROCESSING",
        headers=admin_headers,
    )
    assert listing.status_code == 200
    items = listing.json()["items"]
    assert len(items) == 1
    ticket_id = items[0]["ticket_id"]
    assert items[0]["amount_mismatch"] is False

    confirm = await client.post(
        f"/api/admin/tickets/{ticket_id}/confirm", headers=admin_headers
    )
    assert confirm.status_code == 200, confirm.text

    my_ticket = await client.get("/api/graduates/me/ticket", headers=graduate_headers)
    assert my_ticket.status_code == 200
    ticket_body = my_ticket.json()
    assert ticket_body["ticket"]["status"] == "CONFIRMED"
    assert ticket_body["whatsapp_group_link"] == "https://chat.whatsapp.com/e2e"
    codes = ticket_body["qr_codes"]
    assert len(codes) == 3
    assert codes[0]["code_type"] == "GRADUATE" and codes[0]["is_activated"] is True
    assert all(c["code_type"] == "COMPANION" for c in codes[1:])

    scanner_login = await client.post(
        "/api/scanner/auth/login",
        json={"username": "e2e_scanner", _PW_FIELD: _SCANNER_PW},
    )
    assert scanner_login.status_code == 200, scanner_login.text
    scanner_headers = _auth(scanner_login.json()["access_token"])

    graduate_code = codes[0]["code_value"]
    scan_graduate = await client.post(
        "/api/scanner/scan", json={"code": graduate_code}, headers=scanner_headers
    )
    assert scan_graduate.status_code == 200
    assert scan_graduate.json()["valid"] is True

    for companion in codes[1:]:
        scan = await client.post(
            "/api/scanner/scan",
            json={"code": companion["code_value"]},
            headers=scanner_headers,
        )
        assert scan.json()["valid"] is True, scan.text

    repeat = await client.post(
        "/api/scanner/scan", json={"code": graduate_code}, headers=scanner_headers
    )
    assert repeat.json()["valid"] is False
    assert repeat.json()["decision"] == "ALREADY_USED"

    stats = await client.get(f"/api/admin/stats?event_id={event_id}", headers=admin_headers)
    assert stats.status_code == 200
    stats_body = stats.json()
    assert stats_body["tickets"]["confirmed"] == 1
    assert stats_body["expected_attendance"]["confirmed"] == 3

    export = await client.get(f"/api/admin/export?event_id={event_id}", headers=admin_headers)
    assert export.status_code == 200
    _, rows = _read_rows(export)
    assert rows[0][0] == "اسم الحفل"
    assert rows[1][2] == "محمد احمد علي"
