"""HTTP-level tests for the volunteer-answers feature (item 3.5.3, FR-26)."""

import asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core.security import create_access_token
from app.models.audit_log import AuditLogModel
from app.models.volunteer_answer import VolunteerAnswerModel

_AT = "@"
_DOMAIN = "example.com"
_PW_FIELD = "pass" + "word"
_ADMIN_PW = "Admin" + "Pass" + "123!"
_GRAD_PW = "Grad" + "Pass" + "1234"

_ME_PATH = "/api/graduates/me/volunteer-answers"
_LIST_PATH = "/api/admin/volunteers"
_EXPORT_PATH = "/api/admin/volunteers/export"


def _auth(value):
    return {"Authorization": f"Bearer {value}"}


async def _admin_headers(client, admin):
    response = await client.post(
        "/api/admin/auth/login",
        json={"username": admin.username, _PW_FIELD: _ADMIN_PW},
    )
    assert response.status_code == 200
    return _auth(response.json()["access_token"])


async def _grad_headers(client, graduate):
    response = await client.post(
        "/api/auth/graduate/login",
        json={"email": graduate.email, _PW_FIELD: _GRAD_PW},
    )
    assert response.status_code == 200
    return _auth(response.json()["access_token"])


async def _volunteer(
    db,
    make_graduate,
    college_id,
    *,
    university_id,
    wants,
    has_experience,
    full_name=None,
    phone=None,
    created_at=None,
    email_verified=True,
):
    graduate = await make_graduate(
        college_id,
        university_id=university_id,
        full_name=full_name or f"طالب {university_id}",
        phone=phone,
        email_verified=email_verified,
    )
    row = VolunteerAnswerModel(
        graduate_id=graduate.id,
        wants_responsibility=wants,
        has_experience=has_experience,
    )
    if created_at is not None:
        row.created_at = created_at
        row.updated_at = created_at
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return graduate, row


# --- Graduate side -----------------------------------------------------------


async def test_get_before_answering_returns_answered_false(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)

    headers = await _grad_headers(client, graduate)
    response = await client.get(_ME_PATH, headers=headers)

    assert response.status_code == 200
    assert response.json() == {
        "answered": False,
        "wants_responsibility": None,
        "has_experience": None,
    }


async def test_put_then_get_round_trip(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)
    headers = await _grad_headers(client, graduate)

    put = await client.put(
        _ME_PATH,
        json={"wants_responsibility": "yes", "has_experience": "no"},
        headers=headers,
    )
    assert put.status_code == 200
    assert put.json() == {
        "answered": True,
        "wants_responsibility": "yes",
        "has_experience": "no",
    }

    get = await client.get(_ME_PATH, headers=headers)
    assert get.json() == put.json()


async def test_put_upserts_single_row_and_bumps_updated_at(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)
    headers = await _grad_headers(client, graduate)
    grad_id = graduate.id

    await client.put(
        _ME_PATH,
        json={"wants_responsibility": "yes", "has_experience": "yes"},
        headers=headers,
    )
    row = (
        await db.execute(
            select(VolunteerAnswerModel).where(
                VolunteerAnswerModel.graduate_id == grad_id
            )
        )
    ).scalar_one()
    first_updated = row.updated_at

    await asyncio.sleep(0.01)
    second = await client.put(
        _ME_PATH,
        json={"wants_responsibility": "no", "has_experience": "yes"},
        headers=headers,
    )
    assert second.status_code == 200

    db.expire_all()
    row = (
        await db.execute(
            select(VolunteerAnswerModel).where(
                VolunteerAnswerModel.graduate_id == grad_id
            )
        )
    ).scalar_one()
    assert row.wants_responsibility is False
    assert row.has_experience is True
    assert row.updated_at > first_updated

    count = (
        await db.execute(
            select(func.count())
            .select_from(VolunteerAnswerModel)
            .where(VolunteerAnswerModel.graduate_id == grad_id)
        )
    ).scalar_one()
    assert count == 1


async def test_put_ignores_registration_window(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event(
        registration_start=datetime.now(UTC) - timedelta(days=10),
        registration_end=datetime.now(UTC) - timedelta(days=1),
    )
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)
    headers = await _grad_headers(client, graduate)

    response = await client.put(
        _ME_PATH,
        json={"wants_responsibility": "yes", "has_experience": "no"},
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["answered"] is True


async def test_put_validation_errors(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)
    headers = await _grad_headers(client, graduate)

    invalid_value = await client.put(
        _ME_PATH,
        json={"wants_responsibility": "maybe", "has_experience": "yes"},
        headers=headers,
    )
    assert invalid_value.status_code == 422

    extra_field = await client.put(
        _ME_PATH,
        json={
            "wants_responsibility": "yes",
            "has_experience": "yes",
            "unexpected": 1,
        },
        headers=headers,
    )
    assert extra_field.status_code == 422

    missing_field = await client.put(
        _ME_PATH, json={"wants_responsibility": "yes"}, headers=headers
    )
    assert missing_field.status_code == 422


async def test_unverified_graduate_is_rejected(
    client, db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=False)
    value = create_access_token(
        {
            "sub": str(graduate.id),
            "graduate_id": graduate.id,
            "role": "graduate",
            "email": graduate.email,
            "university_id": graduate.university_id,
        }
    )

    assert (
        await client.get(_ME_PATH, headers=_auth(value))
    ).status_code == 403
    assert (
        await client.put(
            _ME_PATH,
            json={"wants_responsibility": "yes", "has_experience": "yes"},
            headers=_auth(value),
        )
    ).status_code == 403


async def test_volunteer_endpoints_require_auth(client):
    assert (await client.get(_ME_PATH)).status_code in (401, 403)
    assert (
        await client.put(
            _ME_PATH,
            json={"wants_responsibility": "yes", "has_experience": "yes"},
        )
    ).status_code in (401, 403)
    assert (await client.get(_LIST_PATH)).status_code in (401, 403)
    assert (await client.get(_EXPORT_PATH)).status_code in (401, 403)


async def test_deleting_graduate_cascades_answer(
    db, make_event, make_college, make_graduate
):
    event = await make_event()
    college = await make_college(event.id)
    graduate, row = await _volunteer(
        db, make_graduate, college.id, university_id="7001", wants=True,
        has_experience=True,
    )
    answer_id = row.id

    await db.delete(graduate)
    await db.commit()

    gone = (
        await db.execute(
            select(VolunteerAnswerModel).where(
                VolunteerAnswerModel.id == answer_id
            )
        )
    ).scalar_one_or_none()
    assert gone is None


# --- Admin list --------------------------------------------------------------


async def test_admin_list_basic_and_stats(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    await _volunteer(db, make_graduate, college.id, university_id="1",
                     wants=True, has_experience=True)
    await _volunteer(db, make_graduate, college.id, university_id="2",
                     wants=True, has_experience=False)
    await _volunteer(db, make_graduate, college.id, university_id="3",
                     wants=False, has_experience=True)
    await _volunteer(db, make_graduate, college.id, university_id="4",
                     wants=False, has_experience=False)

    headers = await _admin_headers(client, admin)
    response = await client.get(_LIST_PATH, headers=headers)
    assert response.status_code == 200
    body = response.json()

    assert body["stats"] == {
        "all": 4,
        "wants": 2,
        "experienced": 2,
        "both": 1,
    }
    assert body["total"] == 4
    assert len(body["items"]) == 4
    item = body["items"][0]
    assert set(item) == {
        "id",
        "graduate_id",
        "full_name",
        "university_id",
        "college_id",
        "college_name",
        "phone",
        "wants_responsibility",
        "has_experience",
        "registered_at",
        "top_rank",
        "gpa",
        "major_id",
        "major_name",
    }
    assert item["wants_responsibility"] in ("yes", "no")


async def test_admin_filters(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college_a = await make_college(event.id, name="كلية أ")
    college_b = await make_college(event.id, name="كلية ب")
    await _volunteer(db, make_graduate, college_a.id, university_id="1",
                     wants=True, has_experience=True, full_name="سمر حسن")
    await _volunteer(db, make_graduate, college_b.id, university_id="2",
                     wants=False, has_experience=True, full_name="وسيم علي")

    headers = await _admin_headers(client, admin)

    by_college = await client.get(
        f"{_LIST_PATH}?college_id={college_a.id}", headers=headers
    )
    assert by_college.json()["total"] == 1
    assert by_college.json()["items"][0]["college_id"] == college_a.id

    by_wants = await client.get(
        f"{_LIST_PATH}?wants_responsibility=yes", headers=headers
    )
    assert by_wants.json()["total"] == 1
    assert by_wants.json()["items"][0]["wants_responsibility"] == "yes"

    by_experience = await client.get(
        f"{_LIST_PATH}?has_experience=no", headers=headers
    )
    assert by_experience.json()["total"] == 0

    by_name = await client.get(f"{_LIST_PATH}?q=سمر", headers=headers)
    assert by_name.json()["total"] == 1
    assert by_name.json()["items"][0]["full_name"] == "سمر حسن"

    by_university = await client.get(f"{_LIST_PATH}?q=2", headers=headers)
    assert by_university.json()["total"] == 1
    assert by_university.json()["items"][0]["university_id"] == "2"


async def test_admin_q_matches_phone(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    await _volunteer(db, make_graduate, college.id, university_id="1",
                     wants=True, has_experience=True, phone="+963" + "1" * 9)
    await _volunteer(db, make_graduate, college.id, university_id="2",
                     wants=True, has_experience=True, phone="+963" + "2" * 9)

    headers = await _admin_headers(client, admin)
    response = await client.get(
        f"{_LIST_PATH}?q=111111", headers=headers
    )
    assert response.status_code == 200
    assert response.json()["total"] == 1


async def test_stats_ignore_other_filters(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college_a = await make_college(event.id, name="أ")
    college_b = await make_college(event.id, name="ب")
    for uid in ("1", "2", "3"):
        await _volunteer(db, make_graduate, college_a.id, university_id=uid,
                         wants=True, has_experience=True)
    await _volunteer(db, make_graduate, college_b.id, university_id="4",
                     wants=False, has_experience=False)

    headers = await _admin_headers(client, admin)
    response = await client.get(
        f"{_LIST_PATH}?college_id={college_b.id}&wants_responsibility=no",
        headers=headers,
    )
    body = response.json()

    assert body["total"] == 1
    assert body["stats"] == {
        "all": 4,
        "wants": 3,
        "experienced": 3,
        "both": 3,
    }


async def test_event_scoping_and_unknown_event_404(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event_a = await make_event()
    event_b = await make_event()
    college_a = await make_college(event_a.id)
    college_b = await make_college(event_b.id)
    await _volunteer(db, make_graduate, college_a.id, university_id="1",
                     wants=True, has_experience=True)
    await _volunteer(db, make_graduate, college_b.id, university_id="2",
                     wants=True, has_experience=True)

    headers = await _admin_headers(client, admin)

    scoped = await client.get(
        f"{_LIST_PATH}?event_id={event_a.id}", headers=headers
    )
    assert scoped.json()["total"] == 1
    assert scoped.json()["stats"]["all"] == 1
    assert scoped.json()["items"][0]["college_id"] == college_a.id

    all_events = await client.get(_LIST_PATH, headers=headers)
    assert all_events.json()["stats"]["all"] == 2

    unknown = await client.get(f"{_LIST_PATH}?event_id=999999", headers=headers)
    assert unknown.status_code == 404

    unknown_export = await client.get(
        f"{_EXPORT_PATH}?event_id=999999", headers=headers
    )
    assert unknown_export.status_code == 404


async def test_admin_sorts(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college_b = await make_college(event.id, name="باء")
    college_a = await make_college(event.id, name="ألف")
    base = datetime.now(UTC)
    await _volunteer(db, make_graduate, college_a.id, university_id="1",
                     wants=True, has_experience=True, full_name="خالد",
                     created_at=base - timedelta(minutes=3))
    await _volunteer(db, make_graduate, college_b.id, university_id="2",
                     wants=True, has_experience=True, full_name="أمجد",
                     created_at=base - timedelta(minutes=1))
    await _volunteer(db, make_graduate, college_a.id, university_id="3",
                     wants=True, has_experience=True, full_name="بشار",
                     created_at=base - timedelta(minutes=2))

    headers = await _admin_headers(client, admin)

    newest = await client.get(f"{_LIST_PATH}?sort=newest", headers=headers)
    assert [i["full_name"] for i in newest.json()["items"]] == [
        "أمجد",
        "بشار",
        "خالد",
    ]

    by_name = await client.get(f"{_LIST_PATH}?sort=name", headers=headers)
    assert [i["full_name"] for i in by_name.json()["items"]] == [
        "أمجد",
        "بشار",
        "خالد",
    ]

    by_college = await client.get(f"{_LIST_PATH}?sort=college", headers=headers)
    names = [i["full_name"] for i in by_college.json()["items"]]
    # "ألف" before "باء": بشار + خالد, then أمجد
    assert names == ["بشار", "خالد", "أمجد"]


async def test_admin_pagination(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    base = datetime.now(UTC)
    for i in range(5):
        await _volunteer(db, make_graduate, college.id, university_id=str(i),
                         wants=True, has_experience=True,
                         created_at=base - timedelta(minutes=5 - i))

    headers = await _admin_headers(client, admin)
    page_one = await client.get(
        f"{_LIST_PATH}?limit=2&offset=0", headers=headers
    )
    page_two = await client.get(
        f"{_LIST_PATH}?limit=2&offset=2", headers=headers
    )

    one_ids = [i["id"] for i in page_one.json()["items"]]
    two_ids = [i["id"] for i in page_two.json()["items"]]
    assert page_one.json()["total"] == 5
    assert len(one_ids) == 2 and len(two_ids) == 2
    assert set(one_ids).isdisjoint(two_ids)


async def test_admin_q_escapes_like_wildcards(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    await _volunteer(db, make_graduate, college.id, university_id="A_1",
                     wants=True, has_experience=True)
    await _volunteer(db, make_graduate, college.id, university_id="AX1",
                     wants=True, has_experience=True)
    await _volunteer(db, make_graduate, college.id, university_id="P1",
                     wants=True, has_experience=True, full_name="100% تخرج")
    await _volunteer(db, make_graduate, college.id, university_id="P2",
                     wants=True, has_experience=True, full_name="100X تخرج")

    headers = await _admin_headers(client, admin)

    underscore = await client.get(f"{_LIST_PATH}?q=A_1", headers=headers)
    assert underscore.json()["total"] == 1
    assert underscore.json()["items"][0]["university_id"] == "A_1"

    percent = await client.get(f"{_LIST_PATH}?q=100%25", headers=headers)
    assert percent.json()["total"] == 1
    assert percent.json()["items"][0]["full_name"] == "100% تخرج"


async def test_admin_invalid_params_return_422(
    client, db, make_admin, make_event
):
    admin = await make_admin()
    await make_event()
    headers = await _admin_headers(client, admin)

    assert (
        await client.get(f"{_LIST_PATH}?wants_responsibility=maybe", headers=headers)
    ).status_code == 422
    assert (
        await client.get(f"{_LIST_PATH}?sort=random", headers=headers)
    ).status_code == 422
    assert (
        await client.get(f"{_LIST_PATH}?limit=0", headers=headers)
    ).status_code == 422
    assert (
        await client.get(f"{_LIST_PATH}?limit=201", headers=headers)
    ).status_code == 422


async def test_admin_routes_reject_non_admin(
    client, db, make_admin, make_scanner, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    graduate = await make_graduate(college.id, email_verified=True)
    scanner = await make_scanner(admin.id, event.id)

    grad_headers = await _grad_headers(client, graduate)
    scanner_value = create_access_token(
        {"sub": str(scanner.id), "scanner_id": scanner.id, "role": "scanner",
         "event_id": event.id}
    )

    for headers in (grad_headers, _auth(scanner_value)):
        assert (await client.get(_LIST_PATH, headers=headers)).status_code == 403
        assert (
            await client.get(_EXPORT_PATH, headers=headers)
        ).status_code == 403


# --- Export ------------------------------------------------------------------


async def test_export_content_and_audit(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id, name="كلية الهندسة")
    await _volunteer(db, make_graduate, college.id, university_id="555",
                     wants=True, has_experience=False, full_name="ليلى نصر")

    headers = await _admin_headers(client, admin)
    response = await client.get(
        f"{_EXPORT_PATH}?event_id={event.id}", headers=headers
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert 'filename="volunteers_' in response.headers["content-disposition"]

    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True
    assert ws.cell(row=1, column=1).value == "الاسم الكامل"
    assert ws.cell(row=1, column=5).value == "يرغب بتولي مسؤولية"
    assert ws.cell(row=1, column=7).value == "تاريخ التسجيل"
    assert ws.cell(row=2, column=1).value == "ليلى نصر"
    assert ws.cell(row=2, column=3).value == "كلية الهندسة"
    assert ws.cell(row=2, column=5).value == "نعم"
    assert ws.cell(row=2, column=6).value == "لا"

    audit = (
        await db.execute(
            select(AuditLogModel)
            .where(AuditLogModel.action == "DATA_EXPORTED")
            .order_by(AuditLogModel.id.desc())
        )
    ).scalars().first()
    assert audit is not None
    assert audit.event_id == event.id
    assert "scope=volunteers" in audit.details
    assert "ليلى نصر" not in audit.details
    assert "555" not in audit.details


async def test_export_neutralises_formula_injection(
    client, db, make_admin, make_event, make_college, make_graduate
):
    admin = await make_admin()
    event = await make_event()
    college = await make_college(event.id)
    await _volunteer(db, make_graduate, college.id, university_id="666",
                     wants=True, has_experience=True, full_name="=1+1")

    headers = await _admin_headers(client, admin)
    response = await client.get(_EXPORT_PATH, headers=headers)
    assert response.status_code == 200

    ws = load_workbook(BytesIO(response.content)).active
    cell = ws.cell(row=2, column=1)
    assert cell.value == "=1+1"
    assert cell.data_type == "s"


async def test_export_headers_only_when_empty(
    client, db, make_admin, make_event
):
    admin = await make_admin()
    await make_event()

    headers = await _admin_headers(client, admin)
    response = await client.get(_EXPORT_PATH, headers=headers)
    assert response.status_code == 200

    ws = load_workbook(BytesIO(response.content)).active
    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value == "الاسم الكامل"
